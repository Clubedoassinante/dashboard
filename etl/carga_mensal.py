#!/usr/bin/env python3
"""
ETL mensal — Clube do Assinante GZH → Supabase.

Lê a carteira mensal + a planilha de vouchers de um bucket privado do
Supabase Storage, aplica exatamente as mesmas regras já validadas em
process_month.py (regras_negocio.py), transforma CPF em cpf_hash
(HMAC-SHA256 com um pepper secreto que nunca é gravado no banco) e
grava tudo no Postgres do Supabase.

Rodado pelo GitHub Actions (.github/workflows/carga_mensal.yml), nunca
localmente com credenciais reais na sua máquina, para o service_role
key nunca precisar sair do cofre de secrets.

Variáveis de ambiente esperadas (injetadas pelo workflow a partir dos
secrets do repositório):
    SUPABASE_URL
    SUPABASE_SERVICE_ROLE_KEY
    CPF_PEPPER

Uso:
    python3 carga_mensal.py \
        --ano 2026 --mes 9 --ref-date 2026-09-30 \
        --carteira-storage-path carteira/CARTEIRA_CLUBE_SETEMBRO2026.xlsx \
        --vouchers-storage-path vouchers/Jan_a_set_2026.xlsx \
        [--oficial-vouchers 15000] [--oficial-usuarios 7000] [--oficial-frequencia 2.1] \
        [--manter-arquivo-storage]
"""
import argparse
import csv
import datetime
import hashlib
import hmac
import os
import sys
import tempfile
from collections import Counter

import openpyxl
from supabase import create_client

sys.path.insert(0, os.path.dirname(__file__))
from regras_negocio import (
    is_farmacia_excluida, detecta_sentinela_dinamica, faixa_etaria,
    faixa_tenure, safra_semestral, normaliza_cpf, mapeia_plano_familia,
    classifica_situacao, parse_data_pt,
)

BATCH_SIZE = 1000

# Janela móvel de detalhe linha a linha (cpf_hash por assinatura/voucher).
# Decisão de 01/10/2026, revista no mesmo dia (ver Blueprint, bug #8):
# a primeira versão usava um corte fixo de ano calendário (só 2026 em
# diante guardava detalhe cru), mas 8 meses de 2026 já ocuparam 360MB
# -- um ano inteiro (12 meses) estouraria os 500MB do plano gratuito de
# novo antes do ano terminar, e o problema se repetiria todo ano
# seguinte. Trocado por uma janela que anda sozinha: sempre os últimos
# N meses têm detalhe cru, o resto vira só resumo agregado -- sem
# precisar de uma decisão manual nova a cada ano.
JANELA_MESES_DETALHE = 6


def soma_meses(mes_referencia, deslocamento):
    """Desloca uma data 'AAAA-MM-01' em N meses (deslocamento pode ser
    negativo) e devolve no mesmo formato 'AAAA-MM-01'."""
    ano, mes, _ = (int(p) for p in mes_referencia.split("-"))
    total = ano * 12 + (mes - 1) + deslocamento
    novo_ano, novo_mes = divmod(total, 12)
    return f"{novo_ano:04d}-{novo_mes + 1:02d}-01"


def cpf_para_hash(cpf_normalizado, pepper):
    return hmac.new(pepper.encode(), cpf_normalizado.encode(), hashlib.sha256).hexdigest()


def baixar_do_storage(supabase, bucket, path):
    """Baixa um arquivo do Storage para um arquivo temporário local e
    retorna o caminho local. O arquivo original com CPF em texto puro
    nunca toca disco fora deste processo efêmero do GitHub Actions."""
    conteudo = supabase.storage.from_(bucket).download(path)
    sufixo = os.path.splitext(path)[1] or ".xlsx"
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=sufixo)
    tmp.write(conteudo)
    tmp.close()
    return tmp.name


def carrega_carteira(caminho_local, ref_date, pepper, mes_referencia, arquivo_origem):
    wb_scan = openpyxl.load_workbook(caminho_local, read_only=True, data_only=True)
    linhas = list(wb_scan.active.iter_rows(min_row=2, values_only=True))
    wb_scan.close()
    sentinela = detecta_sentinela_dinamica(linhas)

    registros = []
    descartadas_sem_situacao = 0
    chaves_vistas = set()  # (cpf_hash, cdass) -- detecta duplicata exata na planilha
    duplicatas_descartadas = 0
    for row in linhas:
        if row[0] is None and all(v is None for v in row):
            continue
        if row[0] is None:
            descartadas_sem_situacao += 1
            continue
        row_padded = list(row[:10]) + [None] * (10 - len(row[:10]))
        situacao, sigla_p, _nome_p, cdass, _nome, cpf, data_assinatura, idade, _cidade, familia = row_padded
        situ_label = str(situacao).strip().upper() if situacao else "None"
        if situ_label == "PERIODO ENTREGUE":
            situ_label = "ATIVO"

        cpf_norm = normaliza_cpf(cpf)
        if cpf_norm is None:
            continue  # sem CPF válido -> não dá pra rastrear a pessoa, descarta

        cpf_hash = cpf_para_hash(cpf_norm, pepper)
        cdass_str = str(cdass) if cdass is not None else None

        # Linha duplicada na planilha de origem (mesma pessoa + mesma
        # assinatura aparecendo 2x no export) -- descarta a repetição em
        # vez de derrubar a carga inteira na constraint única do banco.
        # Mesma pessoa com CDASS diferente (2+ assinaturas no mês) continua
        # permitida (ver bug #2 no Blueprint).
        chave = (cpf_hash, cdass_str)
        if chave in chaves_vistas:
            duplicatas_descartadas += 1
            continue
        chaves_vistas.add(chave)

        data_assinatura_dt = data_assinatura if isinstance(data_assinatura, datetime.datetime) else None
        mes_entrada = (
            data_assinatura_dt.replace(day=1).date().isoformat()
            if data_assinatura_dt is not None else None
        )

        registros.append({
            "mes_referencia": mes_referencia,
            "cpf_hash": cpf_hash,
            "cdass": cdass_str,
            "situacao_raw": situ_label,
            "situacao": classifica_situacao(situacao),
            "sigla_produto": (str(sigla_p).strip().upper() if sigla_p else "Indefinido"),
            "plano_familia": mapeia_plano_familia(familia),
            "faixa_etaria": faixa_etaria(idade, sentinela),
            "faixa_tenure": faixa_tenure(data_assinatura_dt, ref_date),
            "safra_semestral": safra_semestral(data_assinatura_dt, ref_date),
            "mes_entrada": mes_entrada,
            "arquivo_origem": arquivo_origem,
        })
    print(f"  carteira: {len(registros)} registros com CPF válido "
          f"({descartadas_sem_situacao} linhas descartadas sem situação, "
          f"{duplicatas_descartadas} linhas duplicadas na planilha descartadas)")
    return registros


def _linhas_voucher_xlsx(caminho_local):
    """Formato antigo (jul/2023-abr/2024): xlsx, colunas
    CPF, Marca, Titulo, Descricao, TipoCupom, Data (datetime binário)."""
    wb = openpyxl.load_workbook(caminho_local, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None and all(v is None for v in row):
            continue
        cpf, marca, _titulo, _descricao, tipo_cupom, data = row[:6]
        if not isinstance(data, datetime.datetime):
            continue
        yield cpf, marca, tipo_cupom, data
    wb.close()


def _linhas_voucher_csv(caminho_local):
    """Formato novo (a partir de maio/2024): CSV, colunas
    Nome, CPF, Marca, Título, Descrição do Cupom, Tipo de Cupom,
    Data em texto pt-BR (ex: "31 de mai. de 2024"), Produtos."""
    with open(caminho_local, encoding="utf-8") as f:
        reader = csv.reader(f)
        next(reader, None)  # cabeçalho
        for row in reader:
            if not row or all(not str(c).strip() for c in row):
                continue
            row = (list(row) + [None] * 8)[:8]
            _nome, cpf, marca, _titulo, _descricao, tipo_cupom, data_txt, _produtos = row
            data = parse_data_pt(data_txt)
            if data is None:
                continue
            yield cpf, marca, tipo_cupom, data


def carrega_vouchers(caminho_local, ano, mes, pepper, mes_referencia, arquivo_origem, cpfs_hash_admin):
    extensao = os.path.splitext(caminho_local)[1].lower()
    if extensao == ".csv":
        linhas = _linhas_voucher_csv(caminho_local)
    else:
        linhas = _linhas_voucher_xlsx(caminho_local)

    registros = []
    excl_farmacia = 0
    excl_anomalia = 0
    sem_cpf = 0
    for cpf, marca, tipo_cupom, data in linhas:
        if not (data.year == ano and data.month == mes):
            continue
        cpf_norm = normaliza_cpf(cpf)
        if cpf_norm is None:
            sem_cpf += 1
            continue
        if isinstance(marca, float) and marca.is_integer():
            marca = str(int(marca))
        cpf_hash = cpf_para_hash(cpf_norm, pepper)

        excluido_farmacia = is_farmacia_excluida(marca)
        excluido_anomalia = cpf_hash in cpfs_hash_admin
        if excluido_farmacia:
            excl_farmacia += 1
        if excluido_anomalia:
            excl_anomalia += 1

        registros.append({
            "cpf_hash": cpf_hash,
            "marca": str(marca) if marca is not None else None,
            "tipo_cupom": str(tipo_cupom) if tipo_cupom is not None else None,
            "data_voucher": data.date().isoformat(),
            "mes_referencia": mes_referencia,
            "excluido_farmacia": excluido_farmacia,
            "excluido_anomalia": excluido_anomalia,
            "motivo_exclusao": "conta_administrativa" if excluido_anomalia else None,
            "arquivo_origem": arquivo_origem,
        })
    liquido = len(registros) - excl_farmacia - excl_anomalia
    print(f"  vouchers (formato {extensao or 'xlsx'}): {len(registros)} no mês (bruto pós-CPF), "
          f"excl_farmacia={excl_farmacia}, excl_anomalia={excl_anomalia}, "
          f"liquido={liquido}, sem_cpf_valido={sem_cpf}")
    return registros


def grava_em_lotes(supabase, tabela, registros):
    for i in range(0, len(registros), BATCH_SIZE):
        lote = registros[i:i + BATCH_SIZE]
        supabase.table(tabela).insert(lote).execute()
    print(f"  {tabela}: {len(registros)} linhas gravadas")


def apaga_mes_em_lotes(supabase, tabela, mes_referencia, lote=BATCH_SIZE):
    """Apaga as linhas de um mes_referencia em lotes de `id`, em vez de um
    único DELETE ... WHERE mes_referencia = X.

    Achado no backfill (30/09/2026): conforme as tabelas crescem com o
    histórico acumulado, um DELETE de um mês inteiro de uma vez só pode
    passar do statement_timeout da API do Supabase (erro 57014,
    "canceling statement due to statement timeout") mesmo filtrando por
    um índice. Apagar em lotes pequenos mantém cada operação individual
    rápida o suficiente, não importa o tamanho que a tabela alcance com
    o tempo -- importante também para a carga mensal recorrente, não só
    para o backfill."""
    total = 0
    while True:
        linhas = (
            supabase.table(tabela)
            .select("id")
            .eq("mes_referencia", mes_referencia)
            .limit(lote)
            .execute()
            .data
        )
        if not linhas:
            break
        ids = [r["id"] for r in linhas]
        supabase.table(tabela).delete().in_("id", ids).execute()
        total += len(ids)
    if total:
        print(f"  {tabela}: {total} linhas antigas apagadas (em lotes de {lote})")


def apaga_antes_de_em_lotes(supabase, tabela, cutoff_mes_referencia, lote=BATCH_SIZE):
    """Como apaga_mes_em_lotes, mas apaga tudo ANTES de uma data de
    corte (em vez de igual a um mês específico) -- usado pra manter a
    janela móvel de detalhe cru (ver JANELA_MESES_DETALHE). Também em
    lotes, pelo mesmo motivo do statement_timeout."""
    total = 0
    while True:
        linhas = (
            supabase.table(tabela)
            .select("id")
            .lt("mes_referencia", cutoff_mes_referencia)
            .limit(lote)
            .execute()
            .data
        )
        if not linhas:
            break
        ids = [r["id"] for r in linhas]
        supabase.table(tabela).delete().in_("id", ids).execute()
        total += len(ids)
    if total:
        print(f"  {tabela}: {total} linhas de mês(es) fora da janela (antes de "
              f"{cutoff_mes_referencia}) apagadas")


def mantem_janela_detalhe(supabase, mes_mais_recente, janela_meses=JANELA_MESES_DETALHE):
    """Arquiva (apaga o detalhe cru de) qualquer mês mais antigo que a
    janela móvel permite. Só apaga -- o resumo agregado desses meses já
    foi gravado antes (grava_resumo roda pra TODO mês, sempre, não só
    pros que ficam na janela), então não tem cálculo a refazer."""
    cutoff = soma_meses(mes_mais_recente, -(janela_meses - 1))
    print(f"Mantendo janela de {janela_meses} meses de detalhe cru (corte: antes de {cutoff})...")
    apaga_antes_de_em_lotes(supabase, "carteira_mensal", cutoff)
    apaga_antes_de_em_lotes(supabase, "vouchers_detalhados", cutoff)


DIMENSOES_COMPOSICAO = [
    ("produto", "sigla_produto"),
    ("plano_familia", "plano_familia"),
    ("faixa_etaria", "faixa_etaria"),
    ("faixa_tenure", "faixa_tenure"),
    ("safra_semestral", "safra_semestral"),
]


def calcula_resumo(registros_carteira, registros_vouchers, mes_referencia,
                    oficial_vouchers=None, oficial_usuarios=None, oficial_frequencia=None):
    """Calcula os 3 resumos agregados (kpis, composição, quem-gerou-por-
    dimensão) a partir dos registros já em memória -- sem precisar
    consultar o banco. Mesma lógica das views antigas (002_views.sql),
    agora pré-calculada e gravada em tabela em vez de somada em tempo
    real a cada consulta (ver 008_resumo_mensal.sql)."""
    ativos = [r for r in registros_carteira if r["situacao"] == "Ativa"]
    canceladas = sum(1 for r in registros_carteira if r["situacao"] == "Cancelada")
    novos_assinantes = sum(1 for r in ativos if r["mes_entrada"] == mes_referencia)
    plano_familia_ativo = sum(1 for r in ativos if r["plano_familia"] == "Ativo")
    carteira_ativa = len(ativos)
    plano_familia_ativo_pct = (
        round(plano_familia_ativo / carteira_ativa * 100, 2) if carteira_ativa else None
    )

    liquidos = [v for v in registros_vouchers if not v["excluido_farmacia"] and not v["excluido_anomalia"]]
    vouchers_liquidos = len(liquidos)
    geradores_cpf = {v["cpf_hash"] for v in liquidos}
    usuarios_unicos_geradores = len(geradores_cpf)
    # ^ mesmo set é reaproveitado por quem chama calcula_resumo() para
    # alimentar resumo_cpfs_geradores_mensal (ver grava_cpfs_geradores) --
    # permite contar usuários únicos deduplicados em qualquer período,
    # inclusive um ano inteiro, sem depender da janela móvel de detalhe
    # cru (ver 008_usuarios_unicos_periodo.sql e Blueprint).
    penetracao_pct = (
        round(usuarios_unicos_geradores / carteira_ativa * 100, 2) if carteira_ativa else None
    )

    kpis = {
        "mes_referencia": mes_referencia,
        "carteira_ativa": carteira_ativa,
        "canceladas": canceladas,
        "novos_assinantes": novos_assinantes,
        "plano_familia_ativo_pct": plano_familia_ativo_pct,
        "vouchers_liquidos": vouchers_liquidos,
        "usuarios_unicos_geradores": usuarios_unicos_geradores,
        "penetracao_pct": penetracao_pct,
        "vouchers_gerados_oficial": oficial_vouchers,
        "usuarios_unicos_oficial": oficial_usuarios,
        "frequencia_uso_oficial": oficial_frequencia,
    }

    composicao_rows = []
    quem_gerou_rows = []
    for dimensao, campo in DIMENSOES_COMPOSICAO:
        contagem_total = Counter()
        contagem_geradores = Counter()
        for r in ativos:
            categoria = r[campo]
            contagem_total[categoria] += 1
            if r["cpf_hash"] in geradores_cpf:
                contagem_geradores[categoria] += 1
        for categoria, total in contagem_total.items():
            composicao_rows.append({
                "mes_referencia": mes_referencia, "dimensao": dimensao,
                "categoria": categoria, "qtd": total,
            })
            quem_gerou_rows.append({
                "mes_referencia": mes_referencia, "dimensao": dimensao,
                "categoria": categoria, "carteira_ativa": total,
                "geradores": contagem_geradores.get(categoria, 0),
            })

    return kpis, composicao_rows, quem_gerou_rows, geradores_cpf


def grava_resumo(supabase, mes_referencia, kpis, composicao_rows, quem_gerou_rows):
    supabase.table("resumo_kpis_mensal").upsert(kpis).execute()
    supabase.table("resumo_composicao_mensal").delete().eq("mes_referencia", mes_referencia).execute()
    if composicao_rows:
        supabase.table("resumo_composicao_mensal").insert(composicao_rows).execute()
    supabase.table("resumo_quem_gerou_mensal").delete().eq("mes_referencia", mes_referencia).execute()
    if quem_gerou_rows:
        supabase.table("resumo_quem_gerou_mensal").insert(quem_gerou_rows).execute()
    print(f"  resumo: kpis + {len(composicao_rows)} linhas de composição + "
          f"{len(quem_gerou_rows)} linhas de quem-gerou")


def grava_cpfs_geradores(supabase, mes_referencia, geradores_cpf):
    """Grava o set de cpf_hash de quem gerou voucher líquido no mês, na
    tabela leve que NÃO entra na janela móvel de detalhe cru (fica pra
    sempre, ver 008_usuarios_unicos_periodo.sql) -- é o que permite
    contar usuários únicos deduplicados num ano inteiro depois, sem
    nunca expor cpf_hash ao front-end (só o COUNT(DISTINCT), via
    fn_usuarios_unicos_periodo)."""
    supabase.table("resumo_cpfs_geradores_mensal").delete().eq("mes_referencia", mes_referencia).execute()
    if geradores_cpf:
        linhas = [{"mes_referencia": mes_referencia, "cpf_hash": h} for h in geradores_cpf]
        grava_em_lotes(supabase, "resumo_cpfs_geradores_mensal", linhas)
    else:
        print("  resumo_cpfs_geradores_mensal: 0 linhas (nenhum gerador líquido no mês)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ano", type=int, required=True)
    ap.add_argument("--mes", type=int, required=True)
    ap.add_argument("--ref-date", required=True, help="AAAA-MM-DD, último dia do mês")
    ap.add_argument("--carteira-storage-path", required=True)
    ap.add_argument("--vouchers-storage-path", required=True)
    ap.add_argument("--oficial-vouchers", type=int, default=None)
    ap.add_argument("--oficial-usuarios", type=int, default=None)
    ap.add_argument("--oficial-frequencia", type=float, default=None)
    ap.add_argument("--manter-arquivo-storage", action="store_true",
                     help="por padrão, apaga os arquivos originais do Storage após gravar com sucesso no banco")
    ap.add_argument("--bucket", default="uploads-planilhas")
    args = ap.parse_args()

    supabase_url = os.environ["SUPABASE_URL"]
    service_role_key = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    pepper = os.environ["CPF_PEPPER"]

    supabase = create_client(supabase_url, service_role_key)

    ref_date = datetime.datetime.strptime(args.ref_date, "%Y-%m-%d")
    mes_referencia = f"{args.ano:04d}-{args.mes:02d}-01"
    arquivo_origem_carteira = args.carteira_storage_path
    arquivo_origem_vouchers = args.vouchers_storage_path

    print(f"== Carga de {args.mes:02d}/{args.ano} (mes_referencia={mes_referencia}) ==")

    # contas administrativas conhecidas (ex.: ADMIN C) -- vem do banco,
    # nunca hardcoded aqui, pra nao repetir um CPF real em codigo
    admin_rows = supabase.table("contas_administrativas").select("cpf_hash").execute().data
    cpfs_hash_admin = {r["cpf_hash"] for r in admin_rows}
    print(f"Contas administrativas carregadas do banco: {len(cpfs_hash_admin)}")

    print("Baixando planilhas do Storage...")
    carteira_local = baixar_do_storage(supabase, args.bucket, args.carteira_storage_path)
    vouchers_local = baixar_do_storage(supabase, args.bucket, args.vouchers_storage_path)

    print("Processando carteira...")
    registros_carteira = carrega_carteira(carteira_local, ref_date, pepper, mes_referencia, arquivo_origem_carteira)

    print("Processando vouchers...")
    registros_vouchers = carrega_vouchers(
        vouchers_local, args.ano, args.mes, pepper, mes_referencia,
        arquivo_origem_vouchers, cpfs_hash_admin,
    )

    os.unlink(carteira_local)
    os.unlink(vouchers_local)

    # substitui por completo os dados desse mes (reprocessamento idempotente,
    # igual ao pipeline antigo -- rodar de novo o mesmo mes so troca os dados,
    # nunca duplica)
    print("Calculando e gravando o resumo agregado (alimenta o dashboard)...")
    kpis, composicao_rows, quem_gerou_rows, geradores_cpf = calcula_resumo(
        registros_carteira, registros_vouchers, mes_referencia,
        args.oficial_vouchers, args.oficial_usuarios, args.oficial_frequencia,
    )
    grava_resumo(supabase, mes_referencia, kpis, composicao_rows, quem_gerou_rows)
    grava_cpfs_geradores(supabase, mes_referencia, geradores_cpf)

    print("Limpando dados antigos deste mês (se houver, para reprocessamento seguro)...")
    apaga_mes_em_lotes(supabase, "carteira_mensal", mes_referencia)
    apaga_mes_em_lotes(supabase, "vouchers_detalhados", mes_referencia)

    print("Gravando detalhe cru no banco...")
    grava_em_lotes(supabase, "carteira_mensal", registros_carteira)
    grava_em_lotes(supabase, "vouchers_detalhados", registros_vouchers)

    mantem_janela_detalhe(supabase, mes_referencia)

    if args.oficial_vouchers is not None:
        supabase.table("vouchers_oficial_mensal").upsert({
            "mes_referencia": mes_referencia,
            "vouchers_gerados_oficial": args.oficial_vouchers,
            "usuarios_unicos_oficial": args.oficial_usuarios,
            "frequencia_uso_oficial": args.oficial_frequencia,
            "arquivo_origem": arquivo_origem_vouchers,
        }).execute()
        print("  vouchers_oficial_mensal: atualizado")

    if not args.manter_arquivo_storage:
        supabase.storage.from_(args.bucket).remove([args.carteira_storage_path, args.vouchers_storage_path])
        print("Arquivos originais removidos do Storage (dado já está seguro no banco, como cpf_hash).")

    print(f"\n== Carga de {args.mes:02d}/{args.ano} concluída com sucesso ==")


if __name__ == "__main__":
    main()
